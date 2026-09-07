"""
snapshot_samvat_anchors.py -- Before/After sheet for find_chaitra_shukla_1()
and find_diwali(), the two whole-year Samvat-boundary anchors that sit
outside the festival registry (see plans/tithi_unification/design.md).

Usage:
    # Run BEFORE this plan's Tasks 2-4 change any panchang.py code:
    .venv-linux/bin/python plans/tithi_unification/snapshot_samvat_anchors.py --mode before

    # ... Tasks 2-4 happen here ...

    # Run AFTER, against the same file, to fill After_Date and light up the
    # diff formulas:
    .venv-linux/bin/python plans/tithi_unification/snapshot_samvat_anchors.py --mode after
"""
import argparse
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from panchang import find_chaitra_shukla_1, find_diwali  # noqa: E402

YEARS = [2025, 2026, 2027, 2028]
LAT, LON, TZ_OFFSET = 19.0760, 72.8777, 5.5  # Mumbai -- matches main.py's sample coords
AYANAMSA = "Lahiri"

OUT_PATH = Path(__file__).resolve().parent / "samvat_anchors.xlsx"
HEADERS = ["year", "anchor", "Before_Date", "After_Date", "Days_Diff", "Status"]


def collect_rows() -> list[dict]:
    rows = []
    for year in YEARS:
        cs1 = find_chaitra_shukla_1(year, LAT, LON, TZ_OFFSET, AYANAMSA)
        diwali = find_diwali(year, LAT, LON, TZ_OFFSET, AYANAMSA)
        rows.append({"year": year, "anchor": "chaitra_shukla_1", "date": cs1})
        rows.append({"year": year, "anchor": "diwali", "date": diwali})
    return rows


def write_before(rows: list[dict]) -> None:
    if OUT_PATH.exists():
        print(f"Refusing to overwrite existing {OUT_PATH}.", file=sys.stderr)
        sys.exit(1)
    wb = Workbook()
    ws = wb.active
    ws.title = "Samvat_Anchors"
    ws.append(HEADERS)
    for r, row in enumerate(rows, start=2):
        ws.cell(row=r, column=1).value = row["year"]
        ws.cell(row=r, column=2).value = row["anchor"]
        c = ws.cell(row=r, column=3)
        c.value = row["date"]
        c.number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=5).value = f'=IF($D{r}="","", $D{r}-$C{r})'
        ws.cell(row=r, column=6).value = (
            f'=IF($D{r}="","PENDING",IF($C{r}=$D{r},"SAME","CHANGED"))'
        )
    wb.save(OUT_PATH)
    print(f"Wrote {len(rows)} baseline rows to {OUT_PATH}")


def write_after(rows: list[dict]) -> None:
    if not OUT_PATH.exists():
        print(f"{OUT_PATH} does not exist -- run --mode before first.", file=sys.stderr)
        sys.exit(1)
    wb = load_workbook(OUT_PATH)
    ws = wb.active
    by_key = {(r["year"], r["anchor"]): r["date"] for r in rows}
    for r in range(2, ws.max_row + 1):
        key = (ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value)
        if key in by_key:
            c = ws.cell(row=r, column=4)
            c.value = by_key[key]
            c.number_format = "yyyy-mm-dd"
    wb.save(OUT_PATH)
    print(f"Filled After_Date for {len(rows)} rows in {OUT_PATH}. "
          f"Open it and check Status != SAME against your reference panchang.")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=["before", "after"], required=True)
    args = parser.parse_args()
    rows = collect_rows()
    (write_before if args.mode == "before" else write_after)(rows)


if __name__ == "__main__":
    main()
