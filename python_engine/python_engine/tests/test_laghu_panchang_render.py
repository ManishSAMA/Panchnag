"""Rendering the Laghu Panchang: HTML for print, the PDF, and the Excel workbook.

Uses hand-built LaghuHalf values (the real domain types) so the layout is tested
without the ephemeris; the data itself is covered by test_laghu_panchang_year.
"""

from __future__ import annotations

import calendar
import importlib.util
from datetime import date
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

import hindi_calendar_pdf
from laghu_panchang import COLUMN_HEADERS, LaghuDay, LaghuHalf
from laghu_panchang_render import render_html, write_excel, write_pdf

EDGE = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")
HAS_RENDERER = importlib.util.find_spec("playwright") is not None or EDGE.exists()


def make_day(day: date, *, shubh: str = "-", events: str = "भ. 25:47 से") -> LaghuDay:
    return LaghuDay(
        date=day, shubh=shubh, jain_tithi="9", udaya_tithi="9", vara="शु",
        tithi_end="13:10", nakshatra="चि.", nakshatra_end="17:29", yoga="अ.",
        moon_sign="तुला", events=events,
    )


def make_month(year: int, month: int, *, events: str = "भ. 25:47 से") -> tuple[LaghuHalf, LaghuHalf]:
    last = calendar.monthrange(year, month)[1]

    def half(first: int, end: int) -> LaghuHalf:
        return LaghuHalf(
            samvat="वीर निर्वाण सं. 2553   विक्रम संवत् 2083",
            date_range=f"{first} जनवरी से {end} जनवरी {year} तक।",
            tithi_range="पौष कृष्ण नवमी से पौष शुक्ल सप्तमी तक",
            sun_times="सूर्योदय – 07:19 बजे, सूर्यास्त – 17:31 बजे",
            days=tuple(
                make_day(date(year, month, d), shubh="卐" if d == 2 else "-", events=events)
                for d in range(first, end + 1)
            ),
        )

    return half(1, 15), half(16, last)


def make_year(months=(1, 2)) -> tuple[LaghuHalf, ...]:
    return tuple(half for month in months for half in make_month(2027, month))


def _fake_renderer(html_path: str, pdf_path: str, **_options) -> bool:
    Path(pdf_path).write_bytes(b"%PDF-1.4 stub")
    return True


def _no_renderer(html_path: str, pdf_path: str, **_options) -> bool:
    return False


# ── HTML ─────────────────────────────────────────────────────────────────

def test_html_has_one_page_per_month_with_two_half_tables():
    html = render_html(make_year(), title="लघु पंचांग 2027", location_name="Delhi")

    assert html.count('class="page"') == 2
    assert html.count("<table") == 4
    assert html.count(COLUMN_HEADERS[-1]) == 4


def test_html_carries_every_header_line_and_the_shubh_mark():
    halves = make_year(months=(1,))
    html = render_html(halves, title="लघु पंचांग 2027", location_name="Delhi")

    for half in halves:
        for line in (half.samvat, half.date_range, half.tithi_range, half.sun_times):
            assert line in html
    assert "卐" in html
    assert all(header in html for header in COLUMN_HEADERS)


def test_html_prints_on_the_booklets_8_by_5_inch_page():
    html = render_html(make_year(), title="t", location_name="Delhi")

    assert "size: 8in 5in" in html


def test_html_escapes_cell_text():
    half = make_month(2027, 1)[0]
    risky = LaghuHalf(**{**half.__dict__, "tithi_range": "<b>&</b>"})

    html = render_html((risky,), title="t", location_name="Delhi")

    assert "&lt;b&gt;&amp;&lt;/b&gt;" in html


# ── fitting the page ─────────────────────────────────────────────────────

# A festival-heavy day, as on Mahavir Jayanti or the start of Paryushan.
CROWDED_DAY = (
    "श्री पद्मप्रभ भगवान केवलज्ञान कल्याणक, नवपद ओली पूर्ण (आयंबिल ओली निष्ठापन), "
    "रत्नत्रय व्रत पूर्ण, णमोकार महामंत्र व्रत, भ. 07:07 से 19:21 तक, पंचक 23:52 तक"
)
HAS_PLAYWRIGHT = importlib.util.find_spec("playwright") is not None


def _printed_layout(html: str) -> list[dict]:
    """Per page: how far its tallest half runs past the page, and the table font size."""
    from playwright.sync_api import sync_playwright

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        try:
            page = browser.new_page()
            page.set_content(html)
            page.emulate_media(media="print")
            return page.evaluate("""() => [...document.querySelectorAll('section.page')].map(p => ({
                overflow: Math.max(...[...p.querySelectorAll('.half')].map(h => h.scrollHeight)) - p.clientHeight,
                font: parseFloat(getComputedStyle(p.querySelector('td')).fontSize),
            }))""")
        finally:
            browser.close()


@pytest.mark.skipif(not HAS_PLAYWRIGHT, reason="needs Playwright")
def test_a_festival_heavy_month_still_fits_on_its_page():
    crowded = make_month(2027, 3, events=CROWDED_DAY)

    [layout] = _printed_layout(render_html(crowded, title="t", location_name="Delhi"))

    assert layout["overflow"] <= 0


@pytest.mark.skipif(not HAS_PLAYWRIGHT, reason="needs Playwright")
def test_only_a_crowded_month_is_shrunk_to_fit():
    sparse, crowded = _printed_layout(render_html(
        make_month(2027, 2) + make_month(2027, 3, events=CROWDED_DAY), title="t", location_name="Delhi",
    ))

    assert sparse["font"] == pytest.approx(5.6 * 96 / 72, abs=0.05)  # the 5.6pt base size, in px
    assert crowded["font"] < sparse["font"]


# ── PDF ──────────────────────────────────────────────────────────────────

def test_pdf_is_returned_when_a_renderer_prints_it(tmp_path):
    with patch.object(hindi_calendar_pdf, "_html_to_pdf", _fake_renderer):
        written = write_pdf(make_year(), tmp_path / "laghu.pdf", title="t", location_name="Delhi")

    assert written == str(tmp_path / "laghu.pdf")


def test_html_is_returned_when_no_renderer_is_available(tmp_path):
    with patch.object(hindi_calendar_pdf, "_html_to_pdf", _no_renderer):
        written = write_pdf(make_year(), tmp_path / "laghu.pdf", title="t", location_name="Delhi")

    assert written.endswith(".html")
    assert Path(written).exists()


@pytest.mark.skipif(not HAS_RENDERER, reason="needs Playwright or Microsoft Edge")
def test_the_real_pdf_is_searchable_text_on_8_by_5_inch_pages(tmp_path):
    from pypdf import PdfReader

    written = write_pdf(make_year(), tmp_path / "laghu.pdf", title="t", location_name="Delhi")
    reader = PdfReader(written)

    assert len(reader.pages) == 2
    box = reader.pages[0].mediabox
    assert (round(float(box.width)), round(float(box.height))) == (576, 360)
    text = reader.pages[0].extract_text()
    assert "जनवरी" in text
    assert "13:10" in text


# ── Excel ────────────────────────────────────────────────────────────────

def test_excel_has_a_sheet_per_month_with_both_halves_side_by_side(tmp_path):
    path = write_excel(make_year(), tmp_path / "laghu.xlsx", title="लघु पंचांग 2027", location_name="Delhi")
    workbook = load_workbook(path)

    assert len(workbook.sheetnames) == 2
    january = workbook[workbook.sheetnames[0]]
    values = [cell for row in january.iter_rows(values_only=True) for cell in row if cell is not None]
    assert values.count(COLUMN_HEADERS[-1]) == 2
    assert values.count("तुला") == 31
    assert "1 जनवरी से 15 जनवरी 2027 तक।" in values
    assert "卐" in values


@pytest.mark.skipif(not HAS_PLAYWRIGHT, reason="needs Playwright")
def test_the_longest_pratipada_cell_stays_on_one_line():
    """'आश्वि.कृ. 01' (Ashvin) and 'मार्ग.कृ. 01-2' (Pratipada + kshaya) wrapped onto two lines."""
    from dataclasses import replace

    from playwright.sync_api import sync_playwright

    first, second = make_month(2027, 9)
    # The second is the longest cell there is: a Pratipada that also carries a kshaya.
    days = (replace(first.days[0], jain_tithi="आश्वि.कृ. 01", udaya_tithi="मार्ग.कृ. 01-2"),) + first.days[1:]
    html = render_html((replace(first, days=days), second), title="t", location_name="Delhi")

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        try:
            page = browser.new_page()
            page.set_content(html)
            page.emulate_media(media="print")
            lines = page.evaluate("""() => [2, 3].map(i => {
                const cell = document.querySelector('tbody tr').children[i];
                const range = document.createRange(); range.selectNodeContents(cell);
                return new Set([...range.getClientRects()].map(r => Math.round(r.top))).size;
            })""")
        finally:
            browser.close()

    assert lines == [1, 1]


@pytest.mark.skipif(not HAS_PLAYWRIGHT, reason="needs Playwright")
def test_no_column_header_breaks_inside_a_word():
    """'शुभ दि.' broke as 'शु / भ / दि.' once the shubh column was narrowed."""
    from playwright.sync_api import sync_playwright

    html = render_html(make_month(2027, 1), title="t", location_name="Delhi")
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        try:
            page = browser.new_page()
            page.set_content(html)
            page.emulate_media(media="print")
            broken = page.evaluate(r"""() => [...document.querySelector('thead tr').children].filter(th => {
                const words = th.textContent.trim().split(/\s+/).length;
                const range = document.createRange(); range.selectNodeContents(th);
                const lines = new Set([...range.getClientRects()].map(r => Math.round(r.top))).size;
                return lines > words;
            }).map(th => th.textContent)""")
        finally:
            browser.close()

    assert broken == []
