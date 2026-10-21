"""Laghu Panchang as an export: request parsing, the service, the API and the CLI.

The booklet's content is tested in test_laghu_panchang_year/_render; here the
ephemeris build is replaced by hand-built halves so these stay fast.
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

import hindi_calendar_pdf
import main as cli
import pdf_generation_service
from app import app
from pdf_generation_service import generate_laghu_panchang_export
from request_parsing import LaghuPanchangRequest, parse_laghu_panchang_request
from tests.test_laghu_panchang_render import make_year

DELHI = {"lat": 28.6139, "lon": 77.2090}


def _fake_renderer(html_path: str, pdf_path: str, **_options) -> bool:
    Path(pdf_path).write_bytes(b"%PDF-1.4 stub")
    return True


def _fake_build(**_kwargs):
    return make_year(months=(1,))


# ── request parsing ──────────────────────────────────────────────────────

def test_a_minimal_request_defaults_to_a_pdf_for_the_default_profile():
    parsed = parse_laghu_panchang_request({"year": 2027, **DELHI})

    assert parsed == LaghuPanchangRequest(
        year=2027, city=None, lat=DELHI["lat"], lon=DELHI["lon"], ayanamsa_name="Lahiri",
        profile="shwetambar_murtipujak_tapagachchha", output_format="pdf",
        clock="standard", sunrise_rule="modern",
    )


def test_format_profile_and_city_are_taken_from_the_request():
    parsed = parse_laghu_panchang_request(
        {"year": "2028", "city": "Jaipur", "format": "both", "profile": "shwetambar_terapanthi", "ayanamsa": "Raman"}
    )

    assert (parsed.year, parsed.city, parsed.output_format, parsed.profile, parsed.ayanamsa_name) == (
        2028, "Jaipur", "both", "shwetambar_terapanthi", "Raman",
    )


@pytest.mark.parametrize(
    "payload, message",
    [
        ({**DELHI}, "year"),
        ({"year": 1800, **DELHI}, "year"),
        ({"year": 2027}, "city"),
        ({"year": 2027, "format": "docx", **DELHI}, "format"),
        ({"year": 2027, "profile": "digambar", **DELHI}, "profile"),
    ],
    ids=["no-year", "year-out-of-range", "no-location", "bad-format", "bad-profile"],
)
def test_invalid_requests_are_rejected_with_the_field_named(payload, message):
    with pytest.raises(ValueError, match=message):
        parse_laghu_panchang_request(payload)


# ── service ──────────────────────────────────────────────────────────────

@pytest.mark.parametrize(
    "output_format, suffixes",
    [("pdf", [".pdf"]), ("excel", [".xlsx"]), ("both", [".pdf", ".xlsx"])],
)
def test_the_service_writes_the_requested_formats(tmp_path, output_format, suffixes):
    with patch.object(pdf_generation_service, "build_laghu_year", _fake_build), \
         patch.object(hindi_calendar_pdf, "_html_to_pdf", _fake_renderer):
        result = generate_laghu_panchang_export(
            year=2027, city="Delhi", ayanamsa_name="Lahiri", output_format=output_format,
            profile="shwetambar_murtipujak_tapagachchha", output_dir=str(tmp_path), **DELHI,
        )

    assert [Path(f["path"]).suffix for f in result["files"]] == suffixes
    assert all(Path(f["path"]).exists() for f in result["files"])
    assert all(f["name"].startswith("Laghu_Panchang_2027_Delhi") for f in result["files"])
    assert result["location"]["timezone"] == "Asia/Kolkata"


def test_the_service_builds_for_the_requested_place_profile_and_ayanamsa(tmp_path):
    with patch.object(pdf_generation_service, "build_laghu_year", side_effect=_fake_build) as build, \
         patch.object(hindi_calendar_pdf, "_html_to_pdf", _fake_renderer):
        generate_laghu_panchang_export(
            year=2027, city=None, ayanamsa_name="Raman", output_format="pdf",
            profile="shwetambar_sthanakvasi", output_dir=str(tmp_path), **DELHI,
        )

    build.assert_called_once_with(
        year=2027, lat=DELHI["lat"], lon=DELHI["lon"], tz_name="Asia/Kolkata",
        ayanamsa="Raman", profile="shwetambar_sthanakvasi", clock="standard", sunrise_rule="modern", school="all",
    )


def test_the_service_passes_the_time_settings_to_the_build(tmp_path):
    with patch.object(pdf_generation_service, "build_laghu_year", side_effect=_fake_build) as build,          patch.object(hindi_calendar_pdf, "_html_to_pdf", _fake_renderer):
        generate_laghu_panchang_export(
            year=2027, city=None, ayanamsa_name="Lahiri", output_format="pdf",
            profile="shwetambar_sthanakvasi", output_dir=str(tmp_path),
            clock="daylight", sunrise_rule="traditional", **DELHI,
        )

    assert build.call_args.kwargs["clock"] == "daylight"
    assert build.call_args.kwargs["sunrise_rule"] == "traditional"


def test_the_endpoint_passes_the_time_settings_through():
    with patch("app.generate_laghu_panchang_export", return_value={
        "year": 2027, "ayanamsa": "Lahiri", "profile": "p", "location": {}, "files": [],
    }) as service:
        response = app.test_client().post(
            "/generate-laghu-panchang",
            json={"year": 2027, "clock": "daylight", "sunrise": "traditional", **DELHI},
        )

    assert response.status_code == 200
    assert (service.call_args.kwargs["clock"], service.call_args.kwargs["sunrise_rule"]) == ("daylight", "traditional")


def test_the_cli_accepts_the_time_settings(tmp_path):
    with patch.object(pdf_generation_service, "build_laghu_year", side_effect=_fake_build) as build,          patch.object(hindi_calendar_pdf, "_html_to_pdf", _fake_renderer):
        cli.main([
            "laghu", "--year", "2027", "--lat", "28.6139", "--lon", "77.2090",
            "--clock", "daylight", "--sunrise", "traditional", "--output", str(tmp_path),
        ])

    assert (build.call_args.kwargs["clock"], build.call_args.kwargs["sunrise_rule"]) == ("daylight", "traditional")


# ── API ──────────────────────────────────────────────────────────────────

def test_the_endpoint_returns_download_links_for_each_file(tmp_path):
    client = app.test_client()
    with patch.object(pdf_generation_service, "build_laghu_year", _fake_build), \
         patch.object(hindi_calendar_pdf, "_html_to_pdf", _fake_renderer), \
         patch("tempfile.gettempdir", return_value=str(tmp_path)):
        response = client.post("/generate-laghu-panchang", json={"year": 2027, "format": "both", **DELHI})

    assert response.status_code == 200
    files = response.get_json()["files"]
    assert [f["format"] for f in files] == ["pdf", "xlsx"]
    download = client.get(files[1]["download_url"])
    assert download.status_code == 200
    assert download.data[:2] == b"PK"  # xlsx is a zip


def test_the_endpoint_rejects_a_bad_request():
    response = app.test_client().post("/generate-laghu-panchang", json={"year": 2027, "format": "docx", **DELHI})

    assert response.status_code == 400
    assert "format" in response.get_json()["error"]


# ── CLI ──────────────────────────────────────────────────────────────────

def test_the_cli_writes_the_booklet_to_the_output_directory(tmp_path):
    with patch.object(pdf_generation_service, "build_laghu_year", _fake_build), \
         patch.object(hindi_calendar_pdf, "_html_to_pdf", _fake_renderer):
        cli.main([
            "laghu", "--year", "2027", "--lat", "28.6139", "--lon", "77.2090",
            "--format", "both", "--output", str(tmp_path),
        ])

    written = sorted(p.suffix for p in tmp_path.rglob("Laghu_Panchang_2027*") if p.suffix != ".html")
    assert written == [".pdf", ".xlsx"]
    workbook = load_workbook(next(tmp_path.rglob("*.xlsx")))
    assert len(workbook.sheetnames) == 1
