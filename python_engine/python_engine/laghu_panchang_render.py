"""
laghu_panchang_render.py -- print the Laghu Panchang as PDF (via HTML) and Excel.

The PDF is real Unicode Devanagari text laid out by a browser, so it can be
searched and copied -- unlike the legacy Kruti-font reference booklet, whose
text extracts as Latin gibberish.
"""

from __future__ import annotations

from html import escape
from itertools import groupby
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import hindi_calendar_pdf
from hindi_calendar_pdf import GREG_MONTHS_HINDI
from laghu_panchang import COLUMN_HEADERS, LaghuHalf

# Column widths in percent of a half-page table, in COLUMN_HEADERS order.
_COLUMN_WIDTHS = (4, 5, 9, 9, 4, 8, 7, 8, 6, 12, 28)

_STYLE = """
@page { size: 8in 5in; margin: 0.12in; }
* { box-sizing: border-box; }
body { margin: 0; font-family: 'Nirmala UI', 'Mangal', 'Noto Sans Devanagari', sans-serif;
       color: #111; background: #fff; }
/* Width is fixed, not left to the viewport, so the fit script measures the printed layout. */
.page { display: flex; gap: 0.08in; width: 7.76in; height: 4.74in; page-break-after: always; break-after: page; }
.page:last-child { page-break-after: auto; break-after: auto; }
.half { flex: 1 1 0; min-width: 0; display: flex; flex-direction: column; }
.head { text-align: center; font-size: 6.6pt; line-height: 1.25; font-weight: 600; }
.head .samvat { color: #7a1f1f; }
.head .span { font-size: 7.4pt; font-weight: 700; }
table { width: 100%; border-collapse: collapse; table-layout: fixed; margin-top: 2px; }
th, td { border: 0.5pt solid #555; padding: 0.5pt 1pt; font-size: var(--cell-size, 5.6pt); line-height: 1.15;
         text-align: center; vertical-align: middle; word-wrap: break-word; }
th { background: #f1e3c6; font-weight: 700; font-size: 5.2pt; }
td.events { text-align: left; }
td.shubh { color: #b3001b; font-weight: 700; }
tr.sunday td { background: #fff4f4; }
.place { font-size: 5pt; color: #666; text-align: right; margin-top: 1px; }
"""

# A festival-heavy half (Mahavir Jayanti, Paryushan) can outgrow its fixed page,
# and the overflow would print over the next month. CSS cannot fit text to a box,
# so shrink the cell text of just that page until both halves fit. Playwright and
# Edge both run it before printing.
_FIT_SCRIPT = """
document.querySelectorAll('section.page').forEach(function (page) {
  var halves = Array.prototype.slice.call(page.querySelectorAll('.half'));
  var overflowing = function () {
    return halves.some(function (half) { return half.scrollHeight > page.clientHeight; });
  };
  for (var size = 5.6; overflowing() && size > 3.4; size -= 0.1) {
    page.style.setProperty('--cell-size', (size - 0.1).toFixed(1) + 'pt');
  }
});
"""


def _half_html(half: LaghuHalf, location_name: str) -> str:
    head = (
        f'<div class="head"><div class="samvat">{escape(half.samvat)}</div>'
        f'<div class="span">{escape(half.date_range)}</div>'
        f"<div>{escape(half.tithi_range)}</div>"
        f"<div>{escape(half.sun_times)}</div></div>"
    )
    columns = "".join(f'<col style="width:{width}%">' for width in _COLUMN_WIDTHS)
    headers = "".join(f"<th>{escape(header)}</th>" for header in COLUMN_HEADERS)
    rows = "".join(
        f'<tr class="{"sunday" if day.vara == "र" else ""}">'
        + "".join(
            f'<td class="{_cell_class(index)}">{escape(cell)}</td>'
            for index, cell in enumerate(day.cells)
        )
        + "</tr>"
        for day in half.days
    )
    table = f"<table><colgroup>{columns}</colgroup><thead><tr>{headers}</tr></thead><tbody>{rows}</tbody></table>"
    return f'<div class="half">{head}{table}<div class="place">{escape(location_name)}</div></div>'


def _cell_class(index: int) -> str:
    return {1: "shubh", len(COLUMN_HEADERS) - 1: "events"}.get(index, "")


def _months(halves: tuple[LaghuHalf, ...]) -> list[list[LaghuHalf]]:
    """Group halves into Gregorian months, preserving order."""
    by_month = groupby(halves, key=lambda half: (half.days[0].date.year, half.days[0].date.month))
    return [list(group) for _, group in by_month]


def render_html(halves: tuple[LaghuHalf, ...], *, title: str, location_name: str) -> str:
    """The whole booklet: one 8in x 5in page per month, two half-month tables each."""
    pages = "".join(
        '<section class="page">' + "".join(_half_html(half, location_name) for half in month) + "</section>"
        for month in _months(halves)
    )
    return (
        '<!DOCTYPE html><html lang="hi"><head><meta charset="utf-8">'
        f"<title>{escape(title)}</title><style>{_STYLE}</style></head>"
        f"<body>{pages}<script>{_FIT_SCRIPT}</script></body></html>"
    )


def write_pdf(halves: tuple[LaghuHalf, ...], path: str | Path, *, title: str, location_name: str) -> str:
    """Write the PDF; returns its path, or the HTML's when no browser could print it."""
    pdf_path = Path(path)
    html_path = pdf_path.with_suffix(".html")
    html_path.write_text(render_html(halves, title=title, location_name=location_name), encoding="utf-8")
    if hindi_calendar_pdf._html_to_pdf(str(html_path), str(pdf_path), css_page_size=True):
        return str(pdf_path)
    return str(html_path)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="2C3E50")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=12)
_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_EXCEL_WIDTHS = (5, 6, 11, 11, 5, 10, 8, 10, 7, 12, 45)
_HEADER_LINES = 4
_GAP_COLUMNS = 1


def _write_half(sheet, half: LaghuHalf, first_column: int) -> None:
    last_column = first_column + len(COLUMN_HEADERS) - 1
    for row, line in enumerate((half.samvat, half.date_range, half.tithi_range, half.sun_times), start=1):
        sheet.merge_cells(start_row=row, start_column=first_column, end_row=row, end_column=last_column)
        cell = sheet.cell(row=row, column=first_column, value=line)
        cell.alignment = Alignment(horizontal="center")
        cell.font = _TITLE_FONT if row == 2 else Font(bold=True)

    header_row = _HEADER_LINES + 1
    for offset, header in enumerate(COLUMN_HEADERS):
        cell = sheet.cell(row=header_row, column=first_column + offset, value=header)
        cell.fill, cell.font, cell.border = _HEADER_FILL, _HEADER_FONT, _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(first_column + offset)].width = _EXCEL_WIDTHS[offset]

    for row, day in enumerate(half.days, start=header_row + 1):
        for offset, value in enumerate(day.cells):
            cell = sheet.cell(row=row, column=first_column + offset, value=value)
            cell.border = _BORDER
            cell.alignment = Alignment(
                horizontal="left" if offset == len(COLUMN_HEADERS) - 1 else "center",
                vertical="center", wrap_text=True,
            )


def write_excel(halves: tuple[LaghuHalf, ...], path: str | Path, *, title: str, location_name: str) -> str:
    """One sheet per month, laid out like the printed page: both halves side by side."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = f"{title} — {location_name}"
    for month in _months(halves):
        first_date = month[0].days[0].date
        sheet = workbook.create_sheet(f"{GREG_MONTHS_HINDI[first_date.month - 1]} {first_date.year}")
        for index, half in enumerate(month):
            _write_half(sheet, half, first_column=1 + index * (len(COLUMN_HEADERS) + _GAP_COLUMNS))
        sheet.freeze_panes = sheet.cell(row=_HEADER_LINES + 2, column=1)
    workbook.save(path)
    return str(path)
